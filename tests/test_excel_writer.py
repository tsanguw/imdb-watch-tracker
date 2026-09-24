import io

from openpyxl import load_workbook
from PIL import Image as PILImage

from src.excel_writer import HEADERS, TIER_STYLE, write_excel
from src.models import WatchlistTitle, YouTubeMatch

# Column indices (0-indexed, matching HEADERS) so tests read clearly and
# stay correct if columns are reordered.
POSTER_COL = 0
TITLE_COL = 1
RATING_COL = 3
GENRES_COL = 4
RUNTIME_COL = 5
DIRECTORS_COL = 6
AVAILABILITY_COL = 7
SUBSCRIPTION_COL = 8
RENT_BUY_COL = 9
FREE_AD_COL = 10

CATEGORY_SHEETS = ("Free", "Subscription", "Rent-Buy", "Unavailable")


def _tiny_png_bytes() -> bytes:
    image = PILImage.new("RGB", (10, 15), color=(10, 20, 30))
    buffer = io.BytesIO()
    image.save(buffer, format="PNG")
    return buffer.getvalue()


def _make_titles():
    return [
        # Has all three options — under strict-tier partitioning this
        # should land ONLY on "Free" (free beats subscription beats rent/buy).
        WatchlistTitle(
            imdb_id="tt1",
            title="Has Every Option",
            year=2020,
            title_type="movie",
            tmdb_found=True,
            imdb_rating=8.1,
            genres=["Comedy", "Drama"],
            runtime_minutes=110,
            directors=["Some Director"],
            subscription=["Netflix"],
            rent_buy=["Amazon Video"],
            free_ad_supported=["Tubi"],
            poster_image=_tiny_png_bytes(),
        ),
        WatchlistTitle(
            imdb_id="tt2",
            title="Only On YouTube Unverified",
            year=2019,
            title_type="movie",
            tmdb_found=True,
            youtube_matches=[YouTubeMatch("Full Movie", "https://youtube.com/x", "Some Studio")],
        ),
        # Has both subscription and rent/buy, no free option — should land
        # ONLY on "Subscription" (subscription beats rent/buy).
        WatchlistTitle(
            imdb_id="tt3",
            title="Subscription And Rent Buy Title",
            year=2018,
            title_type="movie",
            tmdb_found=True,
            subscription=["Netflix"],
            rent_buy=["Apple TV Store"],
        ),
        WatchlistTitle(
            imdb_id="tt4",
            title="Rent Buy Only Title",
            year=2017,
            title_type="movie",
            tmdb_found=True,
            rent_buy=["Apple TV Store"],
        ),
        WatchlistTitle(
            imdb_id="tt5",
            title="Found But No Availability",
            year=2016,
            title_type="movie",
            tmdb_found=True,
        ),
        WatchlistTitle(
            imdb_id="tt6",
            title="Totally Unfound Title",
            year=2015,
            title_type="movie",
            tmdb_found=False,
        ),
    ]


def _row_by_title(ws, title_text):
    for row in ws.iter_rows(min_row=2):
        if row[TITLE_COL].value == title_text:
            return row
    return None


def _titles_on(wb, sheet_name):
    return {row[TITLE_COL].value for row in wb[sheet_name].iter_rows(min_row=2)}


def test_write_excel_creates_five_sheets(tmp_path):
    titles = _make_titles()
    output_path = tmp_path / "watchlist.xlsx"
    write_excel(titles, str(output_path))

    wb = load_workbook(output_path)
    assert wb.sheetnames == ["All titles", "Free", "Subscription", "Rent-Buy", "Unavailable"]


def test_write_excel_header_on_every_sheet(tmp_path):
    titles = _make_titles()
    output_path = tmp_path / "watchlist.xlsx"
    write_excel(titles, str(output_path))

    wb = load_workbook(output_path)
    for name in wb.sheetnames:
        ws = wb[name]
        assert [c.value for c in ws[1]] == HEADERS


def test_autofilter_enabled_on_every_sheet(tmp_path):
    titles = _make_titles()
    output_path = tmp_path / "watchlist.xlsx"
    write_excel(titles, str(output_path))

    wb = load_workbook(output_path)
    for name in wb.sheetnames:
        ws = wb[name]
        assert ws.auto_filter.ref is not None


def test_all_titles_sheet_has_every_title(tmp_path):
    titles = _make_titles()
    output_path = tmp_path / "watchlist.xlsx"
    write_excel(titles, str(output_path))

    wb = load_workbook(output_path)
    assert _titles_on(wb, "All titles") == {t.title for t in titles}


def test_category_sheets_are_a_strict_partition(tmp_path):
    """Every title lands on exactly one of the four category sheets —
    never zero, never more than one."""
    titles = _make_titles()
    output_path = tmp_path / "watchlist.xlsx"
    write_excel(titles, str(output_path))

    wb = load_workbook(output_path)

    assert _titles_on(wb, "Free") == {"Has Every Option", "Only On YouTube Unverified"}
    assert _titles_on(wb, "Subscription") == {"Subscription And Rent Buy Title"}
    assert _titles_on(wb, "Rent-Buy") == {"Rent Buy Only Title"}
    assert _titles_on(wb, "Unavailable") == {"Found But No Availability", "Totally Unfound Title"}

    # No title appears on more than one category sheet.
    seen = {}
    for sheet_name in CATEGORY_SHEETS:
        for title in _titles_on(wb, sheet_name):
            assert title not in seen, f"{title!r} appears on both {seen.get(title)} and {sheet_name}"
            seen[title] = sheet_name

    # Every title landed somewhere.
    all_titles = {t.title for t in titles}
    assert set(seen.keys()) == all_titles


def test_rating_genre_and_other_details(tmp_path):
    titles = _make_titles()
    output_path = tmp_path / "watchlist.xlsx"
    write_excel(titles, str(output_path))

    wb = load_workbook(output_path)
    ws = wb["All titles"]

    row = _row_by_title(ws, "Has Every Option")
    assert row[RATING_COL].value == 8.1
    assert isinstance(row[RATING_COL].value, float)
    assert row[GENRES_COL].value == "Comedy, Drama"
    assert row[RUNTIME_COL].value == 110
    assert row[DIRECTORS_COL].value == "Some Director"
    assert row[AVAILABILITY_COL].value == "Free (ad-supported)"

    missing_row = _row_by_title(ws, "Totally Unfound Title")
    assert missing_row[RATING_COL].value == "—"  # em dash placeholder
    assert missing_row[AVAILABILITY_COL].value == "Not found on TMDb"


def test_write_excel_not_found_shows_not_found_label(tmp_path):
    titles = _make_titles()
    output_path = tmp_path / "watchlist.xlsx"
    write_excel(titles, str(output_path))

    wb = load_workbook(output_path)
    ws = wb["All titles"]

    row = _row_by_title(ws, "Totally Unfound Title")
    assert row[SUBSCRIPTION_COL].value == "Not found"
    assert row[RENT_BUY_COL].value == "Not found"
    assert row[FREE_AD_COL].value == "Not found"


def test_write_excel_hyperlink(tmp_path):
    titles = _make_titles()
    output_path = tmp_path / "watchlist.xlsx"
    write_excel(titles, str(output_path))

    wb = load_workbook(output_path)
    ws = wb["All titles"]
    row = _row_by_title(ws, "Has Every Option")
    assert row[TITLE_COL].hyperlink.target == "https://www.imdb.com/title/tt1/"


def test_write_excel_embeds_poster_only_on_sheets_the_title_appears_on(tmp_path):
    titles = _make_titles()
    output_path = tmp_path / "watchlist.xlsx"
    write_excel(titles, str(output_path))

    wb = load_workbook(output_path)
    # "Has Every Option" is the only title with poster_image set, and under
    # strict partitioning it appears on exactly All titles + Free.
    for sheet_name in ("All titles", "Free"):
        ws = wb[sheet_name]
        assert len(ws._images) == 1
        assert ws._images[0].anchor._from.col == POSTER_COL

    for sheet_name in ("Subscription", "Rent-Buy", "Unavailable"):
        assert len(wb[sheet_name]._images) == 0


def test_tier_style_covers_all_six_tiers():
    assert set(TIER_STYLE.keys()) == {1, 2, 3, 4, 5, 6}
