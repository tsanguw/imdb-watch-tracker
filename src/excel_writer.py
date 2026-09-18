"""Builds/overwrites the tracker spreadsheet. No dated snapshots — this
always writes to the same path, overwriting whatever was there before."""
from __future__ import annotations

from pathlib import Path
from typing import List

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from .models import WatchlistTitle

HEADERS = [
    "Title",
    "Year",
    "Subscription Services",
    "Rent/Buy",
    "Free (Ad-Supported)",
    "Free (YouTube - Unverified)",
    "Last Checked",
]

COLUMN_WIDTHS = [32, 8, 28, 28, 28, 45, 20]

NOT_FOUND = "Not found"
NONE_LABEL = "—"  # em dash


def _join(items: List[str]) -> str:
    return ", ".join(items) if items else NONE_LABEL


def _youtube_cell(title: WatchlistTitle) -> str:
    if not title.youtube_matches:
        return NONE_LABEL
    parts = [
        f"unverified - check manually: {m.video_title} ({m.channel_name}) {m.url}"
        for m in title.youtube_matches
    ]
    return "\n".join(parts)


def write_excel(titles: List[WatchlistTitle], output_path: str) -> None:
    Path(output_path).parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "Watchlist"

    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A2"

    for title in titles:
        if not title.tmdb_found:
            subscription = rent_buy = free_ad = NOT_FOUND
        else:
            subscription = _join(title.subscription)
            rent_buy = _join(title.rent_buy)
            free_ad = _join(title.free_ad_supported)

        row = [
            title.title,
            title.year or "",
            subscription,
            rent_buy,
            free_ad,
            _youtube_cell(title),
            title.last_checked or "",
        ]
        ws.append(row)

        row_idx = ws.max_row
        title_cell = ws.cell(row=row_idx, column=1)
        title_cell.hyperlink = f"https://www.imdb.com/title/{title.imdb_id}/"
        title_cell.font = Font(color="0563C1", underline="single")

        for col in range(1, len(HEADERS) + 1):
            ws.cell(row=row_idx, column=col).alignment = Alignment(vertical="top", wrap_text=True)

    for idx, width in enumerate(COLUMN_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width

    wb.save(output_path)
